import calendar as calendar_module
import hashlib
import logging
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta, timezone
from functools import wraps
from logging.handlers import RotatingFileHandler
from zoneinfo import ZoneInfo

import click
from docx import Document
from docx.shared import Pt, RGBColor
from dotenv import load_dotenv
from flask import (
    Flask,
    abort,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from flask_login import (
    LoginManager,
    current_user,
    login_required,
    login_user,
    logout_user,
)
from flask_migrate import Migrate, stamp, upgrade
from sqlalchemy import inspect
from werkzeug.security import check_password_hash, generate_password_hash

from models import Assinatura, Contrato, Imovel, LoginAttempt, Locador, agora_utc, db
from validacao import parse_valor, uf_valida, validar

load_dotenv()

# Chaves de exemplo que já apareceram no repositório: se alguma delas estiver
# em uso, qualquer um consegue forjar a sessão — então nem sobe.
SECRET_KEYS_PUBLICAS = {
    "dev-secret-key-troque-em-producao",
    "troque-por-uma-chave-aleatoria-em-producao",
}
SECRET_KEY = os.environ.get("SECRET_KEY", "")
if not SECRET_KEY or SECRET_KEY in SECRET_KEYS_PUBLICAS:
    sys.exit(
        "SECRET_KEY não definida (ou é a chave de exemplo). Crie o .env com:\n"
        f'  SECRET_KEY={secrets.token_hex(32)}'
    )

app = Flask(__name__, instance_relative_config=True)
app.secret_key = SECRET_KEY
app.config["TEMPLATES_AUTO_RELOAD"] = True
os.makedirs(app.instance_path, exist_ok=True)
app.config["SQLALCHEMY_DATABASE_URI"] = f"sqlite:///{os.path.join(app.instance_path, 'app.db')}"

# Cookies de sessão: fora do alcance de JavaScript e não enviados em requisições
# vindas de outros sites. Com HTTPS, defina COOKIE_SECURE=1 para que só
# trafeguem criptografados.
COOKIE_SECURE = os.environ.get("COOKIE_SECURE") == "1"
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=COOKIE_SECURE,
    REMEMBER_COOKIE_HTTPONLY=True,
    REMEMBER_COOKIE_SAMESITE="Lax",
    REMEMBER_COOKIE_SECURE=COOKIE_SECURE,
)

db.init_app(app)
migrate = Migrate(app, db, render_as_batch=True)

# O banco guarda tudo em UTC; datas mostradas ao usuário e escritas no
# contrato usam o horário de Brasília.
FUSO = ZoneInfo("America/Sao_Paulo")


def hoje_local():
    return datetime.now(FUSO).date()


@app.template_filter("data_hora")
def data_hora(valor):
    # datetime naive em UTC (como sai do banco) -> "dd/mm/aaaa hh:mm" em Brasília
    if not valor:
        return ""
    return valor.replace(tzinfo=timezone.utc).astimezone(FUSO).strftime("%d/%m/%Y %H:%M")


log = logging.getLogger("imofacil")
log.setLevel(logging.INFO)
log.propagate = False
if not log.handlers:
    _formato = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    _console = logging.StreamHandler()
    _console.setFormatter(_formato)
    _arquivo = RotatingFileHandler(
        os.path.join(app.instance_path, "imofacil.log"), maxBytes=1_000_000, backupCount=5, encoding="utf-8"
    )
    _arquivo.setFormatter(_formato)
    log.addHandler(_console)
    log.addHandler(_arquivo)


def ip_cliente():
    return request.remote_addr or "desconhecido"

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "index"

CONTRATOS_DIR = "contratos"
CONTRATOS_GERADOS_DIR = f"{CONTRATOS_DIR}/gerados"
TEMPLATE_PATH = f"{CONTRATOS_DIR}/Contrato_de_Locacao_Residencial.docx"
# Azul da marca, usado nos títulos do modelo de contrato.
COR_TITULO_CONTRATO = RGBColor(0x1E, 0x59, 0x8A)

# Mapeia o nome do campo usado internamente para o placeholder real dentro
# do template .docx (alguns usam acentos, outros não).
FIELD_MAP = {
    # Locatário
    "nome_locatario": "nome_locatario",
    "rg_locatario": "rg_locatario",
    "cpf_locatario": "cpf_locatario",
    "email_locatario": "email_locatario",
    "telefone_locatario": "telefone_locatario",
    "endereco_locatario": "endereço_locatario",
    "nacionalidade_locatario": "nacionalidade_locatario",
    "estado_locatario": "estado_locatario",
    "profissao_locatario": "profissão_locatario",
    # Locador
    "nome_locador": "nome_locador",
    "rg_locador": "rg_locador",
    "cpf_locador": "cpf_locador",
    "email_locador": "email_locador",
    "telefone_locador": "telefone_locador",
    "endereco_locador": "endereço_locador",
    "nacionalidade_locador": "nacionalidade_locador",
    "estado_locador": "estado_locador",
    "profissao_locador": "profissão_locador",
    # Imóvel
    "endereco_imovel": "endereco_imovel",
    "bairro_imovel": "bairro_imovel",
    "cidade_imovel": "cidade_imovel",
    "estado_imovel": "estado_imovel",
    "cep_imovel": "cep_imovel",
    # Termos do contrato
    "duracao_contrato": "duracao_contrato",
    "data_inicio": "data_inicio",
    "data_termino": "data_termino",
    "valor": "valor",
    "forma_pagamento": "forma_pagamento",
    # Data em que o contrato é gerado (preenchimento pelo locatário)
    "dia_assinatura": "dia_assinatura",
    "mes": "mes",
    "ano": "ano",
}

MESES = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]

LOCATARIO_FIELDS = [
    "nome_locatario",
    "rg_locatario",
    "cpf_locatario",
    "email_locatario",
    "telefone_locatario",
    "endereco_locatario",
    "nacionalidade_locatario",
    "estado_locatario",
    "profissao_locatario",
]

PERFIL_FIELDS = [
    "nome",
    "rg",
    "cpf",
    "email",
    "telefone",
    "endereco",
    "nacionalidade",
    "estado_civil",
    "profissao",
]

# Proteção contra força bruta no login: bloqueia por usuário (protege a conta)
# e por IP (protege o servidor de tentar muitos usuários de uma vez).
LOGIN_MAX_TENTATIVAS_USUARIO = 5
LOGIN_MAX_TENTATIVAS_IP = 10
LOGIN_JANELA_BLOQUEIO = timedelta(minutes=15)


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(Locador, int(user_id))


def locador_required(view):
    @wraps(view)
    @login_required
    def wrapped(*args, **kwargs):
        if current_user.role != "locador":
            abort(403)
        if not current_user.ativo:
            logout_user()
            abort(403)
        if not current_user.perfil_completo and request.endpoint != "perfil":
            return redirect(url_for("perfil"))
        return view(*args, **kwargs)

    return wrapped


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped(*args, **kwargs):
        if current_user.role != "admin":
            abort(403)
        return view(*args, **kwargs)

    return wrapped


# Enquanto a senha for a temporária (definida pelo admin), o usuário só pode
# trocar a senha ou sair.
ROTAS_LIBERADAS_SENHA_TEMPORARIA = {"trocar_senha", "logout", "static", "index"}


@app.before_request
def exigir_troca_de_senha():
    if (
        current_user.is_authenticated
        and current_user.senha_temporaria
        and request.endpoint not in ROTAS_LIBERADAS_SENHA_TEMPORARIA
    ):
        return redirect(url_for("trocar_senha"))


def destino_apos_login(user):
    if user.senha_temporaria:
        return url_for("trocar_senha")
    return url_for("admin_locadores" if user.role == "admin" else "dashboard")


def registrar_tentativa_login(usuario, ip, sucesso):
    db.session.add(LoginAttempt(username=usuario, ip=ip, sucesso=sucesso))
    db.session.commit()


def login_bloqueado(usuario, ip):
    limite = agora_utc() - LOGIN_JANELA_BLOQUEIO
    falhas_usuario = LoginAttempt.query.filter(
        LoginAttempt.username == usuario,
        LoginAttempt.sucesso.is_(False),
        LoginAttempt.created_at >= limite,
    ).count()
    falhas_ip = LoginAttempt.query.filter(
        LoginAttempt.ip == ip,
        LoginAttempt.sucesso.is_(False),
        LoginAttempt.created_at >= limite,
    ).count()
    return falhas_usuario >= LOGIN_MAX_TENTATIVAS_USUARIO or falhas_ip >= LOGIN_MAX_TENTATIVAS_IP


def seed_usuarios():
    if Locador.query.first() is None:
        administrador = Locador(
            username="admin",
            password_hash=generate_password_hash("admin123"),
            role="admin",
        )
        locador_teste = Locador(
            username="locador",
            password_hash=generate_password_hash("locador123"),
            role="locador",
        )
        db.session.add_all([administrador, locador_teste])
        db.session.commit()


@app.route("/")
def index():
    if current_user.is_authenticated:
        return redirect(destino_apos_login(current_user))
    return render_template("index.html")


@app.route("/login", methods=["POST"])
def login():
    usuario = request.form.get("usuario", "").strip()
    senha = request.form.get("senha", "")
    ip = ip_cliente()

    if login_bloqueado(usuario, ip):
        log.warning("login bloqueado usuario=%r ip=%s", usuario, ip)
        return jsonify({"error": "Muitas tentativas de login. Aguarde alguns minutos e tente novamente."}), 429

    user = Locador.query.filter_by(username=usuario).first()
    if not user or not check_password_hash(user.password_hash, senha):
        registrar_tentativa_login(usuario, ip, False)
        log.warning("login falhou usuario=%r ip=%s", usuario, ip)
        return jsonify({"error": "Usuário ou senha inválidos."}), 401
    if not user.ativo:
        registrar_tentativa_login(usuario, ip, False)
        log.warning("login de conta desativada usuario=%r ip=%s", usuario, ip)
        return jsonify({"error": "Sua conta está desativada. Fale com o administrador."}), 403

    registrar_tentativa_login(usuario, ip, True)
    login_user(user)
    log.info("login usuario=%r ip=%s", usuario, ip)
    return jsonify({"result": "Login bem-sucedido!", "redirect": destino_apos_login(user)})


@app.route("/logout", methods=["POST"])
@login_required
def logout():
    log.info("logout usuario=%r", current_user.username)
    logout_user()
    return jsonify({"result": "Logout realizado com sucesso!"})


@app.route("/perfil", methods=["GET", "POST"])
@login_required
def perfil():
    if current_user.role == "admin":
        # O administrador não entra em contratos, então não tem dados
        # pessoais — o perfil dele serve só para trocar a senha.
        return redirect(url_for("trocar_senha"))

    if request.method == "GET":
        return render_template("perfil.html", primeiro_acesso=not current_user.perfil_completo)

    dados = {campo: request.form.get(campo, "").strip() for campo in PERFIL_FIELDS}
    missing = [campo for campo, valor in dados.items() if not valor]
    if missing:
        return render_template(
            "perfil.html",
            primeiro_acesso=not current_user.perfil_completo,
            dados=dados,
            error="Preencha todos os campos.",
        ), 400

    erro = validar({"cpf": dados["cpf"], "rg": dados["rg"], "email": dados["email"], "telefone": dados["telefone"]})
    if erro:
        return render_template(
            "perfil.html", primeiro_acesso=not current_user.perfil_completo, dados=dados, error=erro
        ), 400

    nova_senha = request.form.get("nova_senha", "")
    if nova_senha and len(nova_senha) < 6:
        return render_template(
            "perfil.html",
            primeiro_acesso=not current_user.perfil_completo,
            dados=dados,
            error="A nova senha precisa ter pelo menos 6 caracteres.",
        ), 400

    primeiro_acesso = not current_user.perfil_completo
    for campo, valor in dados.items():
        setattr(current_user, campo, valor)
    if nova_senha:
        current_user.password_hash = generate_password_hash(nova_senha)
    current_user.perfil_completo = True
    db.session.commit()
    log.info("perfil atualizado usuario=%r senha_alterada=%s", current_user.username, bool(nova_senha))

    if primeiro_acesso:
        flash("Perfil concluído. Bem-vindo(a)!", "sucesso")
        return redirect(url_for("dashboard"))
    flash("Perfil salvo.", "sucesso")
    return redirect(url_for("perfil"))


@app.route("/trocar-senha", methods=["GET", "POST"])
@login_required
def trocar_senha():
    obrigatoria = current_user.senha_temporaria
    if request.method == "GET":
        return render_template("trocar_senha.html", obrigatoria=obrigatoria)

    nova_senha = request.form.get("nova_senha", "")
    confirmacao = request.form.get("confirmar_senha", "")
    erro = None
    if len(nova_senha) < 6:
        erro = "A nova senha precisa ter pelo menos 6 caracteres."
    elif nova_senha != confirmacao:
        erro = "As senhas não conferem."
    elif check_password_hash(current_user.password_hash, nova_senha):
        erro = "A nova senha precisa ser diferente da atual."
    if erro:
        return render_template("trocar_senha.html", obrigatoria=obrigatoria, error=erro), 400

    current_user.password_hash = generate_password_hash(nova_senha)
    current_user.senha_temporaria = False
    db.session.commit()
    log.info("senha alterada usuario=%r", current_user.username)
    flash("Senha alterada com sucesso.", "sucesso")
    if obrigatoria:
        return redirect(destino_apos_login(current_user))
    return redirect(url_for("trocar_senha"))


@app.route("/dashboard")
@locador_required
def dashboard():
    hoje = hoje_local()
    em_30_dias = hoje + timedelta(days=30)

    total_imoveis = Imovel.query.filter_by(locador_id=current_user.id).count()

    contratos_do_locador = Contrato.query.join(Imovel).filter(Imovel.locador_id == current_user.id)

    alocacoes_realizadas = contratos_do_locador.filter(Contrato.status.in_(["preenchido", "assinado"])).count()
    alocacoes_ativas = contratos_do_locador.filter(
        Contrato.status.in_(["preenchido", "assinado"]),
        Contrato.data_inicio <= hoje,
        Contrato.data_termino >= hoje,
    ).count()
    proximos_30_dias = contratos_do_locador.filter(
        Contrato.data_inicio > hoje,
        Contrato.data_inicio <= em_30_dias,
    ).count()

    return render_template(
        "dashboard.html",
        total_imoveis=total_imoveis,
        alocacoes_realizadas=alocacoes_realizadas,
        alocacoes_ativas=alocacoes_ativas,
        proximos_30_dias=proximos_30_dias,
    )


@app.route("/imoveis")
@locador_required
def imoveis():
    lista = Imovel.query.filter_by(locador_id=current_user.id).order_by(Imovel.created_at.desc()).all()
    return render_template("imoveis_lista.html", imoveis=lista)


CAMPOS_IMOVEL = ["nome", "endereco", "bairro", "cidade", "estado", "cep"]


def dados_imovel_do_form():
    # -> (dados, erro)
    dados = {campo: request.form.get(campo, "").strip() for campo in CAMPOS_IMOVEL}
    if any(not valor for valor in dados.values()):
        return dados, "Preencha todos os campos."
    dados["estado"] = dados["estado"].upper()
    return dados, validar({"uf": dados["estado"], "cep": dados["cep"]})


@app.route("/imoveis/novo", methods=["GET", "POST"])
@locador_required
def novo_imovel():
    if request.method == "GET":
        return render_template("imovel_form.html")

    dados, erro = dados_imovel_do_form()
    if erro:
        return render_template("imovel_form.html", error=erro, dados=dados), 400

    imovel = Imovel(locador_id=current_user.id, **dados)
    db.session.add(imovel)
    db.session.commit()
    log.info("imovel criado id=%s locador=%r", imovel.id, current_user.username)
    flash(f"Imóvel \"{imovel.nome}\" cadastrado.", "sucesso")
    return redirect(url_for("imoveis"))


@app.route("/imoveis/<int:imovel_id>/editar", methods=["GET", "POST"])
@locador_required
def editar_imovel(imovel_id):
    imovel = Imovel.query.filter_by(id=imovel_id, locador_id=current_user.id).first()
    if not imovel:
        abort(404)

    if request.method == "GET":
        return render_template("imovel_form.html", imovel=imovel)

    dados, erro = dados_imovel_do_form()
    if erro:
        return render_template("imovel_form.html", imovel=imovel, dados=dados, error=erro), 400

    for campo, valor in dados.items():
        setattr(imovel, campo, valor)
    db.session.commit()
    log.info("imovel editado id=%s locador=%r", imovel.id, current_user.username)
    flash("Alterações salvas.", "sucesso")
    return redirect(url_for("imoveis"))


@app.route("/imoveis/<int:imovel_id>/excluir", methods=["POST"])
@locador_required
def excluir_imovel(imovel_id):
    imovel = Imovel.query.filter_by(id=imovel_id, locador_id=current_user.id).first()
    if not imovel:
        abort(404)

    if imovel.contratos:
        flash("Não é possível excluir um imóvel com contratos gerados.", "erro")
        return redirect(url_for("imoveis"))

    db.session.delete(imovel)
    db.session.commit()
    log.info("imovel excluido id=%s locador=%r", imovel_id, current_user.username)
    flash(f"Imóvel \"{imovel.nome}\" excluído.", "sucesso")
    return redirect(url_for("imoveis"))


# Filtros da lista de contratos: (status na URL, rótulo). "preenchido" vem
# primeiro porque é o que depende de uma ação do locador.
FILTROS_CONTRATOS = [
    ("", "Todos"),
    ("preenchido", "Falta sua assinatura"),
    ("pendente", "Com o locatário"),
    ("assinado", "Assinados"),
    ("cancelado", "Cancelados"),
]


@app.route("/contratos")
@locador_required
def contratos():
    filtro = request.args.get("status", "")
    if filtro not in dict(FILTROS_CONTRATOS):
        filtro = ""

    todos = (
        Contrato.query.join(Imovel)
        .filter(Imovel.locador_id == current_user.id)
        .order_by(Contrato.created_at.desc())
        .all()
    )
    contagem = {status: sum(1 for c in todos if c.status == status) for status, _ in FILTROS_CONTRATOS if status}
    contagem[""] = len(todos)
    lista = [c for c in todos if not filtro or c.status == filtro]
    return render_template(
        "contratos_lista.html", contratos=lista, filtros=FILTROS_CONTRATOS, filtro=filtro, contagem=contagem
    )


@app.route("/contratos/<int:contrato_id>/visualizar")
@locador_required
def visualizar_contrato(contrato_id):
    contrato = (
        Contrato.query.join(Imovel)
        .filter(Contrato.id == contrato_id, Imovel.locador_id == current_user.id)
        .first()
    )
    if not contrato or contrato.status not in ("preenchido", "assinado"):
        abort(404)

    integridade = integridade_contrato(contrato)
    if integridade == "alterado":
        log.warning("integridade: arquivo do contrato %s não confere com o hash registrado", contrato.id)

    # Lê o .docx gerado (e não o template) — é exatamente o que foi assinado,
    # incluindo o bloco de assinaturas quando o contrato já está completo.
    documento = Document(f"{CONTRATOS_GERADOS_DIR}/{contrato.token}.docx")
    paragrafos = [p.text for p in documento.paragraphs if p.text.strip()]
    assinaturas = sorted(contrato.assinaturas, key=lambda a: a.assinado_em)
    return render_template(
        "contrato_visualizar.html",
        contrato=contrato,
        paragrafos=paragrafos,
        assinaturas=assinaturas,
        integridade=integridade,
    )


@app.route("/contratos/<int:contrato_id>/download")
@locador_required
def baixar_contrato_locador(contrato_id):
    contrato = (
        Contrato.query.join(Imovel)
        .filter(Contrato.id == contrato_id, Imovel.locador_id == current_user.id)
        .first()
    )
    if not contrato or contrato.status not in ("preenchido", "assinado"):
        abort(404)

    try:
        return enviar_pdf(contrato)
    except Exception:
        log.exception("falha ao converter contrato id=%s para PDF", contrato.id)
        flash("Não foi possível gerar o PDF agora. Tente novamente em instantes.", "erro")
        return redirect(url_for("contratos"))


NOMES_MESES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]


@app.route("/calendario")
@locador_required
def calendario():
    hoje = hoje_local()
    ano = request.args.get("ano", type=int) or hoje.year
    mes = request.args.get("mes", type=int) or hoje.month
    if mes < 1 or mes > 12:
        mes = hoje.month

    primeiro_dia = date(ano, mes, 1)
    ultimo_dia = date(ano, mes, calendar_module.monthrange(ano, mes)[1])

    # Assinados ocupam o imóvel (vermelho); preenchidos já têm o aceite do
    # locatário e só esperam a assinatura do locador (amarelo).
    contratos_do_mes = (
        Contrato.query.join(Imovel)
        .filter(
            Imovel.locador_id == current_user.id,
            Contrato.status.in_(["assinado", "preenchido"]),
            Contrato.data_inicio <= ultimo_dia,
            Contrato.data_termino >= primeiro_dia,
        )
        .order_by(Contrato.data_inicio)
        .all()
    )

    ocupacoes = {}
    for contrato in contratos_do_mes:
        dia = max(contrato.data_inicio, primeiro_dia)
        fim = min(contrato.data_termino, ultimo_dia)
        while dia <= fim:
            ocupacoes.setdefault(dia.day, []).append(
                {
                    "id": contrato.id,
                    "imovel": contrato.imovel.nome,
                    "locatario": contrato.nome_locatario,
                    "assinado": contrato.status == "assinado",
                }
            )
            dia += timedelta(days=1)

    semanas = calendar_module.Calendar(firstweekday=6).monthdayscalendar(ano, mes)

    mes_anterior, ano_mes_anterior = (12, ano - 1) if mes == 1 else (mes - 1, ano)
    mes_seguinte, ano_mes_seguinte = (1, ano + 1) if mes == 12 else (mes + 1, ano)

    return render_template(
        "calendario.html",
        semanas=semanas,
        ocupacoes=ocupacoes,
        mes=mes,
        ano=ano,
        nome_mes=NOMES_MESES[mes - 1],
        mes_anterior=mes_anterior,
        ano_mes_anterior=ano_mes_anterior,
        mes_seguinte=mes_seguinte,
        ano_mes_seguinte=ano_mes_seguinte,
        hoje=hoje,
        contratos_do_mes=contratos_do_mes,
        mes_atual=(ano == hoje.year and mes == hoje.month),
    )


@app.route("/contratos/novo", methods=["GET", "POST"])
@locador_required
def novo_contrato():
    imoveis_do_locador = Imovel.query.filter_by(locador_id=current_user.id).order_by(Imovel.nome).all()

    if request.method == "GET":
        return render_template("contrato_novo.html", imoveis=imoveis_do_locador)

    imovel_id = request.form.get("imovel_id", "")
    imovel = next((i for i in imoveis_do_locador if str(i.id) == imovel_id), None)

    campos = ["duracao_contrato", "data_inicio", "data_termino", "valor", "forma_pagamento"]
    dados = {campo: request.form.get(campo, "").strip() for campo in campos}
    missing = [campo for campo, valor in dados.items() if not valor]

    if not imovel or missing:
        return render_template(
            "contrato_novo.html",
            imoveis=imoveis_do_locador,
            error="Selecione um imóvel e preencha todos os termos do contrato.",
        ), 400

    try:
        data_inicio = date.fromisoformat(dados["data_inicio"])
        data_termino = date.fromisoformat(dados["data_termino"])
    except ValueError:
        return render_template(
            "contrato_novo.html", imoveis=imoveis_do_locador, error="Datas inválidas."
        ), 400

    if (data_termino - data_inicio).days <= 1:
        return render_template(
            "contrato_novo.html",
            imoveis=imoveis_do_locador,
            error="A duração do contrato precisa ser maior que 1 dia.",
        ), 400

    valor = parse_valor(dados["valor"])
    if valor is None:
        return render_template(
            "contrato_novo.html",
            imoveis=imoveis_do_locador,
            error="Valor inválido. Use, por exemplo, 1500 ou 1.500,00.",
        ), 400

    contrato = Contrato(
        imovel_id=imovel.id,
        duracao_contrato=dados["duracao_contrato"],
        data_inicio=data_inicio,
        data_termino=data_termino,
        valor=valor,
        forma_pagamento=dados["forma_pagamento"],
    )
    db.session.add(contrato)
    db.session.commit()
    log.info("contrato criado id=%s imovel=%s locador=%r", contrato.id, imovel.id, current_user.username)

    link = url_for("contrato_publico", token=contrato.token, _external=True)
    return render_template("contrato_novo.html", imoveis=imoveis_do_locador, link=link)


@app.route("/contratos/<int:contrato_id>/cancelar", methods=["POST"])
@locador_required
def cancelar_contrato(contrato_id):
    contrato = (
        Contrato.query.join(Imovel)
        .filter(Contrato.id == contrato_id, Imovel.locador_id == current_user.id)
        .first()
    )
    if not contrato:
        abort(404)

    if contrato.status == "pendente":
        contrato.status = "cancelado"
        db.session.commit()
        log.info("contrato cancelado id=%s locador=%r", contrato.id, current_user.username)
        flash("Contrato cancelado.", "sucesso")
    else:
        flash("Só é possível cancelar contratos pendentes.", "erro")

    return redirect(url_for("contratos"))


@app.route("/contratos/<int:contrato_id>/reenviar", methods=["POST"])
@locador_required
def reenviar_contrato(contrato_id):
    contrato = (
        Contrato.query.join(Imovel)
        .filter(Contrato.id == contrato_id, Imovel.locador_id == current_user.id)
        .first()
    )
    if not contrato:
        abort(404)

    if contrato.status == "pendente":
        contrato.token_expires_at = agora_utc() + timedelta(hours=24)
        db.session.commit()
        log.info("link do contrato renovado id=%s locador=%r", contrato.id, current_user.username)
        flash("Link renovado por mais 24 horas. Copie e envie ao locatário.", "sucesso")
    else:
        flash("Só é possível reenviar contratos pendentes.", "erro")

    return redirect(url_for("contratos"))


@app.route("/contratos/<int:contrato_id>/assinar", methods=["GET", "POST"])
@locador_required
def assinar_contrato(contrato_id):
    contrato = (
        Contrato.query.join(Imovel)
        .filter(Contrato.id == contrato_id, Imovel.locador_id == current_user.id)
        .first()
    )
    if not contrato:
        abort(404)

    ja_assinado_pelo_locador = any(a.papel == "locador" for a in contrato.assinaturas)
    if contrato.status != "preenchido" or ja_assinado_pelo_locador:
        flash("Este contrato não está aguardando sua assinatura.", "erro")
        return redirect(url_for("contratos"))

    if request.method == "GET":
        return render_template("contrato_assinar.html", contrato=contrato)

    nome = request.form.get("nome", "").strip()
    if not nome:
        return render_template(
            "contrato_assinar.html", contrato=contrato, error="Digite seu nome completo para assinar."
        ), 400

    caminho = caminho_contrato(contrato)
    contrato.assinaturas.append(
        Assinatura(
            papel="locador",
            nome=nome,
            ip=ip_cliente(),
            user_agent=request.headers.get("User-Agent", "")[:255],
            hash_documento=hash_arquivo(caminho),
        )
    )
    contrato.status = "assinado"
    db.session.flush()

    # Banco e arquivo mudam juntos: o .docx final vai para um arquivo
    # temporário e só substitui o original depois do commit. Se algo falhar
    # no caminho, nada muda — nem o status, nem o arquivo.
    temporario = caminho_temporario()
    try:
        anexar_bloco_assinaturas(contrato, caminho, temporario)
        contrato.hash_final = hash_arquivo(temporario)
        db.session.commit()
        os.replace(temporario, caminho)
    except Exception:
        db.session.rollback()
        remover_se_existir(temporario)
        log.exception("falha ao assinar contrato id=%s", contrato.id)
        return render_template(
            "contrato_assinar.html", contrato=contrato, error="Não foi possível concluir a assinatura. Tente novamente."
        ), 500

    log.info("contrato assinado pelo locador id=%s locador=%r ip=%s", contrato.id, current_user.username, ip_cliente())
    flash("Contrato assinado! O locatário já pode baixar a versão final.", "sucesso")
    return redirect(url_for("contratos"))


@app.route("/contrato/<token>")
def contrato_publico(token):
    contrato = Contrato.query.filter_by(token=token).first()
    if not contrato:
        abort(404)

    if contrato.status == "assinado":
        estado = "assinado"
    elif contrato.status == "preenchido":
        estado = "preenchido"
    elif contrato.status == "cancelado":
        estado = "cancelado"
    elif contrato.link_expirado:
        estado = "expirado"
    else:
        estado = "valido"

    if estado != "valido":
        return render_template("cliente.html", token=token, estado=estado)

    return render_template(
        "cliente.html",
        token=token,
        estado=estado,
        contrato=contrato,
        paragrafos=previa_contrato(contrato),
    )


@app.route("/contrato/<token>/enviar", methods=["POST"])
def enviar_contrato(token):
    contrato = Contrato.query.filter_by(token=token).first()
    if not contrato:
        abort(404)
    if contrato.status == "assinado":
        return jsonify({"error": "Este contrato já foi assinado e não pode mais ser alterado."}), 400
    if contrato.status == "preenchido":
        return jsonify({"error": "Este link já foi utilizado."}), 400
    if contrato.status == "cancelado":
        return jsonify({"error": "Este contrato foi cancelado pelo locador."}), 400
    if contrato.link_expirado:
        return jsonify({"error": "Este link expirou. Peça ao locador para reenviar um novo link."}), 400

    dados_locatario = {campo: request.form.get(campo, "").strip() for campo in LOCATARIO_FIELDS}
    missing = [campo for campo, valor in dados_locatario.items() if not valor]
    if missing:
        return jsonify({"error": f"Campos obrigatórios não preenchidos: {', '.join(missing)}"}), 400

    erro = validar(
        {
            "cpf": dados_locatario["cpf_locatario"],
            "rg": dados_locatario["rg_locatario"],
            "email": dados_locatario["email_locatario"],
            "telefone": dados_locatario["telefone_locatario"],
        }
    )
    if erro:
        return jsonify({"error": erro}), 400

    # Mesmo esquema da assinatura do locador: gera em arquivo temporário e só
    # coloca no lugar definitivo depois que o banco confirmou.
    caminho = caminho_contrato(contrato)
    temporario = caminho_temporario()
    try:
        gerarContratos(contrato, dados_locatario, temporario)

        for campo, valor in dados_locatario.items():
            setattr(contrato, campo, valor)
        contrato.status = "preenchido"
        contrato.filled_at = agora_utc()

        # O checkbox "Aceito os termos e condições" (obrigatório no form) já é o
        # consentimento do locatário — registramos a assinatura dele agora, com
        # o hash do documento que ele acabou de ver/preencher.
        db.session.add(
            Assinatura(
                contrato_id=contrato.id,
                papel="locatario",
                nome=dados_locatario["nome_locatario"],
                documento=dados_locatario["cpf_locatario"],
                ip=ip_cliente(),
                user_agent=request.headers.get("User-Agent", "")[:255],
                hash_documento=hash_arquivo(temporario),
            )
        )
        db.session.commit()
        os.replace(temporario, caminho)
    except FileNotFoundError:
        db.session.rollback()
        remover_se_existir(temporario)
        log.exception("modelo de contrato não encontrado")
        return jsonify({"error": "Modelo de contrato não encontrado no servidor."}), 500
    except Exception:
        db.session.rollback()
        remover_se_existir(temporario)
        log.exception("falha ao gerar contrato id=%s", contrato.id)
        return jsonify({"error": "Não foi possível salvar o contrato gerado."}), 500

    log.info("contrato preenchido pelo locatario id=%s ip=%s", contrato.id, ip_cliente())
    return jsonify({"result": "Contrato gerado com sucesso!"})


# O link público expõe CPF/RG/endereço do locatário — por isso o download por
# ele só fica disponível por alguns dias depois da última assinatura. O
# locador continua baixando pelo painel (baixar_contrato_locador).
DOWNLOAD_PUBLICO_PRAZO = timedelta(days=7)


@app.route("/contrato/<token>/download")
def download_contrato(token):
    contrato = Contrato.query.filter_by(token=token).first()
    if not contrato or contrato.status not in ("preenchido", "assinado"):
        abort(404)

    ultima = contrato.ultima_assinatura_em
    if not ultima or agora_utc() - ultima > DOWNLOAD_PUBLICO_PRAZO:
        return render_template("cliente.html", token=token, estado="download_expirado"), 410

    try:
        resposta = enviar_pdf(contrato)
    except Exception:
        log.exception("falha ao converter contrato id=%s para PDF", contrato.id)
        return render_template("cliente.html", token=token, estado="download_erro"), 503
    log.info("download publico contrato id=%s ip=%s", contrato.id, ip_cliente())
    return resposta


@app.route("/admin")
@admin_required
def admin_locadores():
    locadores = Locador.query.filter_by(role="locador").order_by(Locador.username).all()
    return render_template("admin_locadores.html", locadores=locadores)


@app.route("/admin/locadores/novo", methods=["GET", "POST"])
@admin_required
def admin_novo_locador():
    if request.method == "GET":
        return render_template("admin_locador_form.html")

    usuario = request.form.get("usuario", "").strip()
    senha = request.form.get("senha", "")

    if not usuario or not senha:
        return render_template(
            "admin_locador_form.html", error="Preencha usuário e senha.", usuario=usuario
        ), 400

    if Locador.query.filter_by(username=usuario).first():
        return render_template(
            "admin_locador_form.html", error="Já existe um usuário com esse nome.", usuario=usuario
        ), 400

    if len(senha) < 6:
        return render_template(
            "admin_locador_form.html", error="A senha precisa ter pelo menos 6 caracteres.", usuario=usuario
        ), 400

    # A senha criada pelo admin é provisória: o locador troca no 1º login.
    locador = Locador(
        username=usuario, password_hash=generate_password_hash(senha), role="locador", senha_temporaria=True
    )
    db.session.add(locador)
    db.session.commit()
    log.info("admin %r criou locador %r", current_user.username, usuario)
    flash(f"Locador \"{usuario}\" criado. Ele vai definir uma nova senha no primeiro acesso.", "sucesso")
    return redirect(url_for("admin_locadores"))


@app.route("/admin/locadores/<int:locador_id>/alternar-acesso", methods=["POST"])
@admin_required
def admin_alternar_acesso(locador_id):
    locador = Locador.query.filter_by(id=locador_id, role="locador").first()
    if not locador:
        abort(404)

    locador.ativo = not locador.ativo
    db.session.commit()
    acao = "reativado" if locador.ativo else "desativado"
    log.info("admin %r: locador %r %s", current_user.username, locador.username, acao)
    flash(f"Acesso de \"{locador.username}\" {acao}.", "sucesso")
    return redirect(url_for("admin_locadores"))


def gerar_senha_temporaria():
    # 10 caracteres sem os ambíguos (0/O, 1/l/I) — fácil de ditar por telefone.
    alfabeto = "abcdefghijkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    return "".join(secrets.choice(alfabeto) for _ in range(10))


@app.route("/admin/locadores/<int:locador_id>/redefinir-senha", methods=["POST"])
@admin_required
def admin_redefinir_senha(locador_id):
    locador = Locador.query.filter_by(id=locador_id, role="locador").first()
    if not locador:
        abort(404)

    senha = gerar_senha_temporaria()
    locador.password_hash = generate_password_hash(senha)
    locador.senha_temporaria = True
    db.session.commit()
    log.info("admin %r redefiniu a senha do locador %r", current_user.username, locador.username)
    # Mostrada uma única vez: não fica guardada em lugar nenhum.
    flash(
        f"Senha temporária de \"{locador.username}\": {senha} — repasse ao locador; "
        "ele vai precisar trocá-la no próximo login.",
        "sucesso",
    )
    return redirect(url_for("admin_locadores"))


@app.route("/admin/tentativas")
@admin_required
def admin_tentativas():
    tentativas = (
        LoginAttempt.query.filter_by(sucesso=False).order_by(LoginAttempt.created_at.desc()).limit(50).all()
    )
    return render_template("admin_tentativas.html", tentativas=tentativas)


def hash_arquivo(caminho):
    with open(caminho, "rb") as arquivo:
        return hashlib.sha256(arquivo.read()).hexdigest()


def caminho_contrato(contrato):
    return f"{CONTRATOS_GERADOS_DIR}/{contrato.token}.docx"


SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def converter_para_pdf(origem, destino):
    # LibreOffice em modo headless. Cada conversão usa um perfil próprio numa
    # pasta temporária: com o perfil padrão, duas conversões ao mesmo tempo
    # (dois downloads simultâneos) travam uma à outra.
    if not SOFFICE:
        raise RuntimeError("LibreOffice (soffice) não encontrado no servidor")
    with tempfile.TemporaryDirectory() as pasta:
        subprocess.run(
            [
                SOFFICE,
                f"-env:UserInstallation=file://{pasta}/perfil",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                pasta,
                os.path.abspath(origem),
            ],
            check=True,
            capture_output=True,
            timeout=90,
        )
        gerado = os.path.join(pasta, os.path.splitext(os.path.basename(origem))[0] + ".pdf")
        temporario = caminho_temporario()
        shutil.move(gerado, temporario)
    os.replace(temporario, destino)


def pdf_do_contrato(contrato):
    # O .docx continua sendo o documento oficial (é dele o hash das
    # assinaturas); o PDF é uma cópia para download, refeita sempre que o
    # .docx for mais novo — por exemplo, depois que o locador assina.
    docx = caminho_contrato(contrato)
    pdf = os.path.splitext(docx)[0] + ".pdf"
    if not os.path.exists(pdf) or os.path.getmtime(pdf) < os.path.getmtime(docx):
        converter_para_pdf(docx, pdf)
    return pdf


def enviar_pdf(contrato):
    partes = ["Contrato", contrato.imovel.nome, contrato.nome_locatario or ""]
    nome = " - ".join(re.sub(r'[\\/:*?"<>|]+', "", parte).strip() for parte in partes if parte.strip())
    return send_file(
        os.path.abspath(pdf_do_contrato(contrato)),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"{nome}.pdf",
    )


def caminho_temporario():
    # Na mesma pasta do destino, para o os.replace ser atômico.
    os.makedirs(CONTRATOS_GERADOS_DIR, exist_ok=True)
    descritor, caminho = tempfile.mkstemp(suffix=".docx.tmp", dir=CONTRATOS_GERADOS_DIR)
    os.close(descritor)
    return caminho


def remover_se_existir(caminho):
    try:
        os.remove(caminho)
    except FileNotFoundError:
        pass


def integridade_contrato(contrato):
    # Compara o .docx em disco com o hash gravado no momento da última
    # assinatura. -> "integro", "alterado" ou "indisponivel" (sem hash
    # registrado ou arquivo ausente).
    if contrato.status == "assinado":
        esperado = contrato.hash_final
    else:
        esperado = next((a.hash_documento for a in contrato.assinaturas if a.papel == "locatario"), None)
    caminho = caminho_contrato(contrato)
    if not esperado or not os.path.exists(caminho):
        return "indisponivel"
    return "integro" if hash_arquivo(caminho) == esperado else "alterado"


def anexar_bloco_assinaturas(contrato, origem, destino):
    doc = Document(origem)
    titulo = doc.add_paragraph()
    titulo.paragraph_format.space_before = Pt(24)
    titulo.paragraph_format.keep_with_next = True
    run = titulo.add_run("ASSINATURAS ELETRÔNICAS")
    run.bold = True
    run.font.color.rgb = COR_TITULO_CONTRATO
    for assinatura in sorted(contrato.assinaturas, key=lambda a: a.assinado_em):
        papel = "Locatário" if assinatura.papel == "locatario" else "Locador"
        linha = doc.add_paragraph()
        linha.paragraph_format.space_after = Pt(4)
        rotulo = linha.add_run(f"{papel}: ")
        rotulo.bold = True
        rotulo.font.size = Pt(9)
        texto = assinatura.nome
        if assinatura.documento:
            texto += f" (doc. {assinatura.documento})"
        texto += f" — assinado em {data_hora(assinatura.assinado_em)} (horário de Brasília) pelo IP {assinatura.ip}"
        linha.add_run(texto).font.size = Pt(9)
    doc.save(destino)


def dados_contrato(contrato, dados_locatario):
    imovel = contrato.imovel
    locador = imovel.locador

    dados = dict(dados_locatario)
    dados["nome_locador"] = locador.nome
    dados["rg_locador"] = locador.rg
    dados["cpf_locador"] = locador.cpf
    dados["email_locador"] = locador.email
    dados["telefone_locador"] = locador.telefone
    dados["endereco_locador"] = locador.endereco
    dados["nacionalidade_locador"] = locador.nacionalidade
    dados["estado_locador"] = locador.estado_civil
    dados["profissao_locador"] = locador.profissao
    dados["endereco_imovel"] = imovel.endereco
    dados["bairro_imovel"] = imovel.bairro
    dados["cidade_imovel"] = imovel.cidade
    dados["estado_imovel"] = imovel.estado
    dados["cep_imovel"] = imovel.cep
    dados["duracao_contrato"] = contrato.duracao_contrato
    dados["data_inicio"] = contrato.data_inicio.strftime("%d/%m/%Y")
    dados["data_termino"] = contrato.data_termino.strftime("%d/%m/%Y")
    dados["valor"] = contrato.valor_formatado
    dados["forma_pagamento"] = contrato.forma_pagamento
    hoje = hoje_local()
    dados["dia_assinatura"] = str(hoje.day)
    dados["mes"] = MESES[hoje.month - 1]
    dados["ano"] = str(hoje.year)
    return dados


def previa_contrato(contrato):
    # Texto do template com tudo que já se sabe (locador, imóvel, termos)
    # preenchido, para o locatário ler antes de assinar. Cada parágrafo vira
    # uma lista de trechos: (texto, None) ou (None, campo_do_locatario) —
    # esses campos são preenchidos na tela conforme ele digita.
    dados = dados_contrato(contrato, {})
    campo_por_placeholder = {placeholder: campo for campo, placeholder in FIELD_MAP.items()}

    paragrafos = []
    for paragraph in Document(TEMPLATE_PATH).paragraphs:
        if not paragraph.text.strip():
            continue
        trechos = []
        for i, parte in enumerate(re.split(r"\{\{(.+?)\}\}", paragraph.text)):
            if i % 2 == 0:
                trechos.append((parte, None))
                continue
            campo = campo_por_placeholder.get(parte)
            if campo in LOCATARIO_FIELDS:
                trechos.append((None, campo))
            elif campo:
                trechos.append((dados[campo] or "", None))
            else:
                # Placeholder que o gerarContratos também não preenche — mostramos
                # igual ao que vai sair no .docx.
                trechos.append(("{{" + parte + "}}", None))
        paragrafos.append(trechos)
    return paragrafos


def gerarContratos(contrato, dados_locatario, destino):
    dados = dados_contrato(contrato, dados_locatario)

    valores = {"{{" + FIELD_MAP[campo] + "}}": valor or "" for campo, valor in dados.items()}
    doc = Document(TEMPLATE_PATH)
    for paragraph in doc.paragraphs:
        if "{{" in paragraph.text:
            substituir_no_paragrafo(paragraph, valores)

    doc.save(destino)


def substituir_no_paragrafo(paragraph, valores):
    # Troca os placeholders numa única passada, da esquerda para a direita:
    # o texto já inserido nunca é reprocessado (se o locatário digitar
    # "{{valor}}" no nome, fica escrito assim mesmo).
    #
    # Mexe só nos "runs" (trechos com formatação) que cada placeholder
    # ocupa — atribuir paragraph.text apagaria o negrito de "LOCADOR:" etc.
    # O Word às vezes quebra um mesmo {{placeholder}} em vários runs: o valor
    # entra no primeiro e o resto do placeholder é removido dos seguintes.
    runs = paragraph.runs
    padrao = re.compile("|".join(re.escape(placeholder) for placeholder in valores))
    busca_a_partir = 0
    while True:
        texto = "".join(run.text for run in runs)
        achado = padrao.search(texto, busca_a_partir)
        if not achado:
            return
        inicio, fim = achado.span()
        valor = valores[achado.group()]
        posicao = 0
        for run in runs:
            run_inicio, run_fim = posicao, posicao + len(run.text)
            posicao = run_fim
            if run_fim <= inicio or run_inicio >= fim:
                continue
            corte_inicio = max(inicio, run_inicio) - run_inicio
            corte_fim = min(fim, run_fim) - run_inicio
            novo = valor if run_inicio <= inicio < run_fim else ""
            run.text = run.text[:corte_inicio] + novo + run.text[corte_fim:]
        busca_a_partir = inicio + len(valor)


# Revisão que corresponde ao esquema criado pelo antigo db.create_all().
REVISAO_ESQUEMA_INICIAL = "0001"


def preparar_banco():
    # Aplica as migrações pendentes. Bancos criados antes do Flask-Migrate
    # (sem a tabela alembic_version) já têm o esquema da 0001: são marcados
    # nela antes de subir, para não recriar tabelas que já existem.
    tabelas = set(inspect(db.engine).get_table_names())
    if tabelas and "alembic_version" not in tabelas:
        log.info("banco existente sem controle de versão: marcando na revisão %s", REVISAO_ESQUEMA_INICIAL)
        stamp(revision=REVISAO_ESQUEMA_INICIAL)
    upgrade()


BACKUPS_DIR = "backups"


@app.cli.command("backup")
def backup_comando():
    """Copia o banco e os contratos gerados para backups/<data-hora>/."""
    pasta = os.path.join(BACKUPS_DIR, datetime.now(FUSO).strftime("%Y-%m-%d_%H%M%S"))
    os.makedirs(pasta, exist_ok=True)

    # A API de backup do SQLite copia um instantâneo consistente mesmo com a
    # aplicação no ar (copiar o arquivo direto pode pegar uma escrita no meio).
    origem = sqlite3.connect(os.path.join(app.instance_path, "app.db"))
    copia = sqlite3.connect(os.path.join(pasta, "app.db"))
    with copia:
        origem.backup(copia)
    origem.close()
    copia.close()

    if os.path.isdir(CONTRATOS_GERADOS_DIR):
        shutil.make_archive(os.path.join(pasta, "contratos_gerados"), "zip", CONTRATOS_GERADOS_DIR)

    manter = int(os.environ.get("BACKUP_MANTER", "14"))
    antigos = sorted(os.listdir(BACKUPS_DIR))[:-manter] if manter > 0 else []
    for nome in antigos:
        shutil.rmtree(os.path.join(BACKUPS_DIR, nome), ignore_errors=True)

    log.info("backup criado em %s (%d antigo(s) removido(s))", pasta, len(antigos))
    click.echo(pasta)


@app.cli.command("redefinir-senha")
@click.argument("usuario")
def redefinir_senha_comando(usuario):
    """Gera uma senha temporária para USUARIO (útil se o admin esquecer a dele)."""
    user = Locador.query.filter_by(username=usuario).first()
    if not user:
        raise click.ClickException(f"Usuário {usuario!r} não encontrado.")
    senha = gerar_senha_temporaria()
    user.password_hash = generate_password_hash(senha)
    user.senha_temporaria = True
    db.session.commit()
    log.info("senha de %r redefinida pela linha de comando", usuario)
    click.echo(f"Senha temporária de {usuario}: {senha}")


# Os comandos "flask db ..." gerenciam o banco por conta própria; nos demais
# casos (servidor, backup, redefinir-senha) o banco sobe já atualizado.
if not (os.path.basename(sys.argv[0]) == "flask" and "db" in sys.argv[1:]):
    with app.app_context():
        preparar_banco()
        seed_usuarios()


if __name__ == "__main__":
    app.run(host="0.0.0.0", debug=os.environ.get("FLASK_DEBUG") == "1")
